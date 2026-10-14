"""Reprise des soldes antérieurs : import Excel et saisie manuelle (V5-12)."""
import uuid
from pathlib import Path

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse

from administration.models import School, SchoolYear
from pedagogie.models import Classe, Enrollment, Student
from .models import FeeType, Invoice
from .services.arrears import (add_prior_balance, import_arrears, last_archivable_start, match_students,
                               parse_workbook, student_arrears)
from .services.grouped_payment import parse_amount


def _school():
    year = SchoolYear.objects.filter(is_active=True).select_related("school").first()
    return year.school if year else School.objects.filter(is_active=True).first()


def _tmp_dir():
    path = Path(settings.KIS_DATA_DIR) / "tmp_imports"
    path.mkdir(parents=True, exist_ok=True)
    # Ménage des prévisualisations abandonnées (plus d'un jour).
    import time
    for old in path.glob("*.xlsx"):
        try:
            if time.time() - old.stat().st_mtime > 86400:
                old.unlink()
        except OSError:
            pass
    return path


def _remove(path):
    try:
        path.unlink()
    except OSError:
        pass


def _archive_summary(school):
    rows = []
    for year in SchoolYear.objects.filter(school=school, is_archive=True).order_by("-start_date"):
        invoices = Invoice.objects.filter(enrollment__school_year=year, fee_type__category=FeeType.Category.ARREARS)
        total = sum((i.total_amount for i in invoices), 0)
        paid = sum((i.paid_amount for i in invoices), 0)
        rows.append({"year": year, "count": invoices.count(), "total": total, "paid": paid, "rest": total - paid})
    return rows


@login_required
def arrears_import(request):
    school = _school()
    report = None
    token = request.POST.get("token") or ""
    if request.method == "POST":
        try:
            if request.POST.get("confirm") and token:
                path = _tmp_dir() / f"{Path(token).name}.xlsx"
                if not path.exists():
                    raise ValidationError("Le fichier de prévisualisation a expiré : chargez-le de nouveau.")
                with path.open("rb") as fh:
                    report = match_students(parse_workbook(fh), school)
                choices = {r.index: request.POST.get(f"match_{r.index}", r.default_choice) for r in report.importable}
                result = import_arrears(report, school, choices)
                _remove(path)
                text = (f"{result['invoices']} reliquat(s) repris pour un total de {result['total']:,.0f} {school.currency}"
                        f" ; {result['students']} ancien(s) élève(s) créé(s)").replace(",", " ")
                if result["skipped"]:
                    text += f" ; {result['skipped']} ligne(s) ignorée(s)"
                if result["already"]:
                    text += f" ; {result['already']} déjà présent(s)"
                messages.success(request, text + ".")
                return redirect("arrears_import")
            upload = request.FILES.get("file")
            if not upload:
                raise ValidationError("Choisissez le fichier Excel.")
            if not upload.name.lower().endswith(".xlsx"):
                raise ValidationError("Seuls les fichiers .xlsx sont acceptés (Excel 2007 et plus).")
            token = uuid.uuid4().hex
            path = _tmp_dir() / f"{token}.xlsx"
            with path.open("wb") as fh:
                for chunk in upload.chunks():
                    fh.write(chunk)
            with path.open("rb") as fh:
                report = match_students(parse_workbook(fh), school)
        except ValidationError as exc:
            messages.error(request, " ".join(exc.messages))
            report = None
    shown = [r for r in report.rows if r.status != "soldé"] if report else []
    return render(request, "finance/arrears/import.html", {
        "report": report, "shown": shown, "token": token, "school": school, "archives": _archive_summary(school),
        "settled_count": len(report.rows) - len(shown) if report else 0,
    })


def _year_choices(school):
    last = last_archivable_start(school)
    return [(str(y), f"{y}-{y + 1}") for y in range(last, last - 10, -1)] + [("nc", "Antérieur — année non précisée")]


def _start_year(raw, school):
    raw = (raw or "").strip()
    if raw == "nc":
        return None
    if raw not in {value for value, _ in _year_choices(school)}:
        raise ValidationError("Choisissez l'année concernée.")
    return int(raw)


@login_required
def prior_balance_add(request, student_id):
    student = get_object_or_404(Student, pk=student_id)
    school = _school()
    error = None
    if request.method == "POST":
        try:
            start_year = _start_year(request.POST.get("start_year"), school)
            amount = parse_amount(request.POST.get("amount"))
            invoice, created = add_prior_balance(student, school, start_year, amount, request.POST.get("label") or None)
            if created:
                messages.success(request, f"« {invoice.fee_type.name} » ajouté à la situation de {student}.")
            else:
                messages.warning(request, "Ce reliquat existait déjà : il n'a pas été ajouté une seconde fois.")
            return redirect("student_detail", student_id=student.id)
        except ValidationError as exc:
            error = " ".join(exc.messages)
    return render(request, "finance/arrears/add.html", {
        "student": student, "years": _year_choices(school), "error": error, "post": request.POST, "school": school,
    })


@login_required
def arrears_quick(request):
    """Saisie rapide par classe : pour les écoles qui tiennent leurs reliquats sur cahier."""
    school = _school()
    active = SchoolYear.objects.filter(school=school, is_active=True).first()
    classes = (Classe.objects.filter(school_year=active).select_related("education_level", "pedagogical_track")
               .order_by("education_level__order", "group", "name")) if active else Classe.objects.none()
    classe = classes.filter(pk=request.GET.get("classe") or request.POST.get("classe") or 0).first()
    year_raw = request.GET.get("annee") or request.POST.get("annee") or ""
    rows, error = [], None
    if classe and year_raw:
        try:
            start_year = _start_year(year_raw, school)
        except ValidationError as exc:
            error, start_year, classe = " ".join(exc.messages), None, None
    if classe and year_raw and not error:
        enrollments = classe.enrollments.select_related("student").order_by("student__last_name", "student__first_name")
        for e in enrollments:
            rows.append({"student": e.student, "existing": student_arrears(e.student).get(start_year),
                         "posted": request.POST.get(f"amount_{e.student_id}", "")})
        if request.method == "POST":
            added = total = 0
            try:
                with transaction.atomic():
                    for row in rows:
                        if row["existing"]:
                            continue
                        amount = parse_amount(request.POST.get(f"amount_{row['student'].pk}"))
                        if amount:
                            add_prior_balance(row["student"], school, start_year, amount)
                            added += 1
                            total += amount
            except ValidationError as exc:
                error = " ".join(exc.messages)
            else:
                messages.success(request, f"{added} reliquat(s) enregistré(s) pour {total:,.0f} {school.currency}.".replace(",", " "))
                return redirect(f"{reverse('arrears_quick')}?classe={classe.pk}&annee={year_raw}")
    return render(request, "finance/arrears/quick.html", {
        "school": school, "classes": classes, "classe": classe, "years": _year_choices(school),
        "year_raw": year_raw, "rows": rows, "error": error, "post": request.POST,
    })


@login_required
def arrears_template(request):
    """Modèle Excel de reprise ; avec ?eleves=1 il est pré-rempli avec les élèves de l'année en cours."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    school = _school()
    wb = Workbook()
    ws = wb.active
    ws.title = "Reliquats"
    headers = ["Matricule", "Prénom", "Nom", "Classe", "Année", "Reste dû"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F2D5C")
        cell.alignment = Alignment(horizontal="center")
    for col, width in zip("ABCDEF", (14, 22, 18, 16, 14, 14)):
        ws.column_dimensions[col].width = width
    if request.GET.get("eleves") == "1":
        active = SchoolYear.objects.filter(school=school, is_active=True).first()
        if active:
            for e in (Enrollment.objects.filter(school_year=active)
                      .select_related("student", "classe__education_level", "classe__pedagogical_track")
                      .order_by("classe__education_level__order", "classe__group", "student__last_name", "student__first_name")):
                ws.append([e.student.matricule or "", e.student.first_name, e.student.last_name, e.classe.display_name, "", None])
    else:
        last = last_archivable_start(school)
        ws.append(["", "Awa", "TRAORÉ", "6e A", f"{last}-{last + 1}", 25000])
        ws.append(["", "Awa", "TRAORÉ", "5e A", f"{last - 1}-{last}", 10000])
        ws.append(["", "Moussa", "KONÉ", "", "", 40000])
    years = ",".join(label for value, label in _year_choices(school)[:-1])
    validation = DataValidation(type="list", formula1=f'"{years}"', allow_blank=True)
    validation.error = "Écrivez l'année sous la forme 2024-2025, ou laissez vide si elle n'est pas connue."
    ws.add_data_validation(validation)
    validation.add("E2:E2000")
    notes = wb.create_sheet("Mode d'emploi")
    for line in (
        "Une ligne par reliquat : un élève qui doit sur deux années a deux lignes.",
        "Année : 2024-2025 par exemple ; laissez vide si l'année n'est pas connue (« année non précisée »).",
        "Reste dû : montant encore à payer pour cette année, sans espace ni « FCFA ». Les lignes vides ou à 0 sont ignorées.",
        "Matricule : facultatif ; il permet de retrouver l'élève sans ambiguïté.",
        "Les élèves qui ne sont plus à l'école sont créés comme anciens élèves.",
    ):
        notes.append([line])
    notes.column_dimensions["A"].width = 110
    from io import BytesIO
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    name = "modele_reliquats_eleves.xlsx" if request.GET.get("eleves") == "1" else "modele_reliquats.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{name}"'
    return response
