"""Reprise des soldes des années antérieures (Feedbacks V5, point 12).

Un établissement qui adopte KIS Kalanso en cours de route a souvent des
élèves qui doivent encore de l'argent sur les années passées, y compris des
élèves qui ont quitté l'école mais continuent de payer. Pour ne pas repartir
de zéro :

- chaque année antérieure est créée comme « année reprise » (clôturée, sans
  référentiel ni bulletin) ;
- l'élève y reçoit une inscription technique dans la classe « Reprise des
  soldes » et une facture « Reliquat de scolarité AAAA-AAAA » du montant restant
  dû ;
- ces reliquats s'encaissent ensuite comme n'importe quel frais (reçu unique
  avec les frais de l'année en cours, famille comprise).

Un reliquat peut aussi être « antérieur, année non précisée » quand l'école ne
connaît que le total dû : il est alors rangé dans la dernière année reprise et
considéré comme la dette la plus ancienne.

L'import est idempotent : réimporter le même fichier ne crée pas de doublons.
"""
import datetime
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.db import transaction

from administration.models import SchoolYear
from finance.models import FeeType, Invoice
from pedagogie.models import Classe, EducationLevel, Enrollment, Student

ARCHIVE_CLASS_NAME = "Reprise des soldes"
UNSPECIFIED_KEY = "anterieur-non-precise"
UNSPECIFIED_LABEL = "Arriérés antérieurs (année non précisée)"


def _norm(value):
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode("ascii").lower()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())


def first_managed_year(school):
    return SchoolYear.objects.filter(school=school, is_archive=False).order_by("start_date").first()


def last_archivable_start(school):
    """Dernière année que l'on peut reprendre : celle qui précède la première année gérée dans KIS."""
    first = first_managed_year(school)
    return (first.start_date.year if first else datetime.date.today().year) - 1


def archive_year(school, start_year):
    """Année antérieure de reprise (créée au besoin). Refuse une année déjà suivie dans KIS."""
    name = f"{start_year}-{start_year + 1}"
    year = SchoolYear.objects.filter(school=school, name=name).first()
    if year:
        if not year.is_archive:
            raise ValidationError(
                f"L'année {name} est déjà gérée dans KIS Kalanso : ses impayés y sont déjà suivis, rien à reprendre."
            )
        return year
    if start_year > last_archivable_start(school):
        first = first_managed_year(school)
        raise ValidationError(
            f"L'année {name} n'est pas antérieure à la première année gérée dans KIS ({first.name if first else '—'})."
        )
    return SchoolYear.objects.create(
        school=school, start_date=datetime.date(start_year, 10, 1), end_date=datetime.date(start_year + 1, 6, 30),
        is_archive=True, is_closed=True, is_active=False,
    )


def archive_class(year):
    classe = Classe.objects.filter(school_year=year, school=year.school).first()
    if classe:
        return classe
    types = year.school.enabled_school_types()
    level = EducationLevel.objects.filter(school_type__in=types).order_by("order").first() or EducationLevel.objects.order_by("order").first()
    classe = Classe(school=year.school, school_year=year, education_level=level, group="", name=ARCHIVE_CLASS_NAME)
    classe.save()
    return classe


def arrears_fee(year, label=None, unspecified=False):
    if unspecified:
        label = UNSPECIFIED_LABEL
    label = (label or f"Reliquat de scolarité {year.name}").strip()[:100]
    fee, _ = FeeType.objects.get_or_create(
        name=label, school_year=year, classe=None, education_level=None,
        defaults={"category": FeeType.Category.ARREARS, "amount": Decimal("0"), "is_required": False,
                  "cycle_key": UNSPECIFIED_KEY if unspecified else ""},
    )
    return fee


def is_unspecified(invoice):
    return invoice.fee_type.cycle_key == UNSPECIFIED_KEY and invoice.fee_type.category == FeeType.Category.ARREARS


@transaction.atomic
def add_prior_balance(student, school, start_year, amount, label=None):
    """Ajoute (ou retrouve) le reliquat d'un élève pour une année antérieure.

    ``start_year`` vaut ``None`` pour un reliquat dont l'année n'est pas connue.
    Retourne ``(facture, créée)`` ; un reliquat déjà présent n'est jamais dupliqué.
    """
    if amount is None or amount <= 0:
        raise ValidationError("Le reliquat doit être supérieur à zéro.")
    unspecified = start_year in (None, "", 0)
    year = archive_year(school, last_archivable_start(school) if unspecified else int(start_year))
    classe = archive_class(year)
    enrollment = Enrollment.objects.filter(student=student, school_year=year).first()
    if not enrollment:
        enrollment = Enrollment.objects.create(student=student, classe=classe, school_year=year)
    fee = arrears_fee(year, label, unspecified=unspecified)
    invoice = Invoice.objects.filter(enrollment=enrollment, fee_type=fee).first()
    if invoice:
        return invoice, False
    return Invoice.objects.create(enrollment=enrollment, fee_type=fee, total_amount=amount), True


def student_arrears(student):
    """Reliquats déjà enregistrés pour un élève : {année de début ou None: facture}."""
    result = {}
    for inv in Invoice.objects.filter(enrollment__student=student, fee_type__category=FeeType.Category.ARREARS).select_related("fee_type", "enrollment__school_year"):
        result[None if is_unspecified(inv) else inv.enrollment.school_year.start_date.year] = inv
    return result


# --------------------------------------------------------------------------
# Import Excel : tableau propre à l'établissement ou modèle KIS
# --------------------------------------------------------------------------
YEAR_RE = re.compile(r"(20\d{2}|19\d{2})\s*[-/–]\s*(20\d{2}|19\d{2})")


def _year_from(text):
    found = YEAR_RE.search(str(text or ""))
    if found and int(found.group(2)) == int(found.group(1)) + 1:
        return int(found.group(1))
    return None


@dataclass
class ArrearsRow:
    index: int
    line_label: str
    start_year: int = None        # None : année non précisée
    first_name: str = ""
    last_name: str = ""
    matricule: str = ""
    annual: Decimal = None
    paid: Decimal = None
    balance: Decimal = None
    class_label: str = ""
    status: str = ""              # importable : « élève existant », « à vérifier », « nouvel élève »
    message: str = ""
    student: object = None        # rattachement proposé
    candidates: list = field(default_factory=list)

    @property
    def year_label(self):
        return f"{self.start_year}-{self.start_year + 1}" if self.start_year else "Non précisée"

    @property
    def importable(self):
        return self.status in ("élève existant", "à vérifier", "nouvel élève")

    @property
    def default_choice(self):
        if self.student:
            return str(self.student.pk)
        # Homonymes sans choix : on n'invente pas, la ligne attend une décision.
        return "skip" if self.status == "à vérifier" else "new"


@dataclass
class ArrearsReport:
    rows: list = field(default_factory=list)

    @property
    def importable(self):
        return [r for r in self.rows if r.importable]

    @property
    def to_check(self):
        return [r for r in self.rows if r.status == "à vérifier"]

    @property
    def total(self):
        return sum((r.balance for r in self.importable), Decimal("0"))

    @property
    def years(self):
        return sorted({r.year_label for r in self.importable})


def _amount(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).quantize(Decimal("1"))
    text = re.sub(r"[^\d,.\-]", "", str(value)).replace(",", ".")
    if text in ("", "-", "."):
        return None
    try:
        return Decimal(text).quantize(Decimal("1"))
    except InvalidOperation:
        return None


HEADER_KEYS = {
    "first_name": ("prenom", "prenoms"),
    "last_name": ("nom", "noms", "nom de famille"),
    "matricule": ("matricule", "mle", "n matricule"),
    "annual": ("mt annuel", "montant annuel", "annuel", "total annuel", "montant du", "montant total"),
    "paid": ("mt percu", "montant percu", "percu", "total percu", "total paye", "deja paye", "paye"),
    "balance": ("reliquat", "reste", "reste du", "reste a payer", "solde", "credit", "arrieres", "dette"),
    "class_label": ("classe",),
    "year": ("annee", "annee scolaire"),
}


def _header_map(cells):
    mapping = {}
    for idx, cell in enumerate(cells):
        key = _norm(cell)
        if not key:
            continue
        for target, names in HEADER_KEYS.items():
            if target in mapping:
                continue
            if key in names or (target in ("annual", "paid", "balance") and any(key.startswith(n) for n in names)):
                mapping[target] = idx
                break
    if "last_name" in mapping and ("balance" in mapping or ("annual" in mapping and "paid" in mapping)):
        return mapping
    return None


CLASS_RE = re.compile(r"^[A-Za-z0-9ÀÂÉÈÊËÎÏÔÛÙÜÇ°ème\s\-]{1,12}$")


def parse_workbook(file):
    """Lit toutes les feuilles du classeur.

    - Une ligne d'en-têtes (PRENOMS, NOM, Mt ANNUEL, Mt perçu, Reliquat… ou le
      modèle KIS) fixe les colonnes ; elle reste valable pour les blocs suivants
      même s'ils n'ont qu'un titre.
    - Un titre contenant « AAAA - AAAA » fixe l'année des lignes qui suivent ; une
      colonne « Année » sur la ligne prime ; sans aucune année, le reliquat est
      « année non précisée ».
    - Les valeurs calculées (formules, liens vers d'autres classeurs) sont lues
      telles qu'Excel les a enregistrées.
    """
    from openpyxl import load_workbook

    try:
        wb = load_workbook(file, data_only=True, read_only=True)
    except Exception:
        raise ValidationError("Le fichier n'est pas un classeur Excel (.xlsx) lisible.")
    report = ArrearsReport()
    for ws in wb.worksheets:
        current_year = _year_from(ws.title)
        mapping = None
        for line, row in enumerate(ws.iter_rows(values_only=True), 1):
            cells = list(row)
            values = [c for c in cells if c not in (None, "")]
            if not values:
                continue
            header = _header_map(cells)
            if header:
                mapping = header
                current_year = _year_from(" ".join(str(c) for c in values)) or current_year
                continue
            has_number = any(isinstance(c, (int, float, Decimal)) and not isinstance(c, bool) for c in cells)
            title_year = _year_from(" ".join(str(c) for c in values))
            if title_year and not has_number:
                current_year = title_year
                continue
            if not mapping:
                continue
            get = lambda key: cells[mapping[key]] if key in mapping and mapping[key] < len(cells) else None
            last = " ".join(str(get("last_name") or "").split())
            first = " ".join(str(get("first_name") or "").split())
            if not last and not first:
                continue  # ligne de totaux sans nom
            if _norm(last) in ("total", "totaux", "total general") or _norm(first) in ("total", "totaux", "total general"):
                continue
            item = ArrearsRow(index=len(report.rows), line_label=f"{ws.title} · ligne {line}",
                              first_name=first, last_name=last,
                              matricule=str(get("matricule") or "").strip(),
                              annual=_amount(get("annual")), paid=_amount(get("paid")), balance=_amount(get("balance")),
                              class_label=str(get("class_label") or "").strip())
            if "year" in mapping:
                raw = get("year")
                item.start_year = _year_from(raw) if raw not in (None, "") else None
                if raw not in (None, "") and item.start_year is None:
                    item.status, item.message = "erreur", f"Année illisible : « {raw} » (écrivez par ex. 2024-2025)."
            else:
                item.start_year = current_year
            if not item.class_label:
                # Classe notée dans une colonne sans titre, à droite du tableau (ex. « MS », « GS »).
                last_col = max(mapping.values())
                for extra in cells[last_col + 1:]:
                    if isinstance(extra, str) and extra.strip() and CLASS_RE.match(extra.strip()):
                        item.class_label = extra.strip()
                        break
            if item.balance is None and item.annual is not None:
                item.balance = item.annual - (item.paid or Decimal("0"))
            if item.status:
                pass
            elif item.balance is None:
                item.status, item.message = "erreur", "Reliquat illisible."
            elif item.balance <= 0:
                item.status, item.message = "soldé", "Rien à reprendre."
            report.rows.append(item)
    if not report.rows:
        raise ValidationError(
            "Aucun tableau reconnu. Il faut une ligne d'en-têtes avec au moins NOM et Reliquat "
            "(ou Mt ANNUEL et Mt perçu). Vous pouvez aussi partir du modèle KIS."
        )
    return report


def _student_label(student, current_classes):
    label = f"{student.last_name} {student.first_name}"
    if student.matricule:
        label += f" · {student.matricule}"
    classe = current_classes.get(student.pk)
    return label + (f" · {classe}" if classe else " · ancien élève")


def match_students(report, school):
    """Propose un élève pour chaque ligne.

    - « élève existant » : même matricule, ou mêmes nom et prénom (un seul élève) ;
    - « à vérifier » : homonymes, ou nom identique et prénom proche (« Adama F » /
      « Adama Fousseyni ») : la liste « Rattacher à… » permet de choisir ;
    - « nouvel élève » : personne d'approchant, un ancien élève sera créé.
    """
    students = list(Student.objects.all().only("id", "first_name", "last_name", "matricule"))
    by_matricule = {s.matricule.strip().lower(): s for s in students if s.matricule}
    by_last = {}
    for s in students:
        by_last.setdefault(_norm(s.last_name), []).append(s)
    current = first_managed_year(school)
    active = SchoolYear.objects.filter(school=school, is_active=True).first()
    current_classes = {}
    if active:
        for e in Enrollment.objects.filter(school_year=active).select_related("classe__education_level", "classe__pedagogical_track"):
            current_classes[e.student_id] = e.classe.display_name
    limit = last_archivable_start(school)
    seen = {}
    for r in report.rows:
        if r.status:
            continue
        if r.start_year and r.start_year > limit:
            r.status, r.message = "erreur", f"L'année {r.year_label} est déjà gérée dans KIS ({current.name if current else ''} et suivantes)."
            continue
        key = (r.start_year, _norm(r.last_name), _norm(r.first_name), r.matricule.lower())
        if key in seen:
            r.status, r.message = "doublon", f"Même élève et même année qu'à la {seen[key]}."
            continue
        seen[key] = r.line_label
        exact = by_matricule.get(r.matricule.lower()) if r.matricule else None
        same_last = by_last.get(_norm(r.last_name), [])
        first_norm = _norm(r.first_name)
        first_token = first_norm.split(" ")[0] if first_norm else ""
        same_name = [s for s in same_last if _norm(s.first_name) == first_norm]
        close = [s for s in same_last if s not in same_name and first_token and (
            _norm(s.first_name).split(" ")[0] == first_token
            or _norm(s.first_name).startswith(first_norm) or first_norm.startswith(_norm(s.first_name)))]
        if not exact and len(same_name) == 1:
            exact = same_name[0]
        candidates = ([exact] if exact else []) + [s for s in same_name + close if s is not exact]
        r.candidates = [(str(s.pk), _student_label(s, current_classes)) for s in candidates[:10]]
        if exact:
            r.student = exact
            r.status, r.message = "élève existant", "Rattaché automatiquement."
        elif len(same_name) > 1:
            r.status, r.message = "à vérifier", "Plusieurs élèves portent ce nom : choisissez le bon."
            r.student = None
        elif close:
            r.student = close[0] if len(close) == 1 else None
            r.status, r.message = "à vérifier", "Prénom proche d'un élève existant : vérifiez le rattachement."
        else:
            r.status, r.message = "nouvel élève", "Aucun élève approchant : un ancien élève sera créé."
        if r.student:
            existing = student_arrears(r.student).get(r.start_year)
            if existing:
                r.status, r.message = "déjà importé", f"Reliquat {r.year_label} déjà présent pour {r.student}."
    return report


@transaction.atomic
def import_arrears(report, school, choices=None):
    """Enregistre les reliquats. ``choices`` : {index de ligne: id élève | "new" | "skip"}."""
    choices = choices or {}
    by_id = {}
    created = {}  # un ancien élève présent sur plusieurs années du fichier n'est créé qu'une fois
    done = {"invoices": 0, "students": 0, "total": Decimal("0"), "skipped": 0, "already": 0}
    for r in report.importable:
        choice = str(choices.get(r.index, r.default_choice))
        if choice == "skip":
            done["skipped"] += 1
            continue
        if choice == "new":
            key = r.matricule.lower() or (_norm(r.last_name), _norm(r.first_name))
            student = created.get(key)
            if not student:
                student = Student.objects.create(first_name=r.first_name or "—", last_name=r.last_name or "—",
                                                 matricule=r.matricule or None)
                created[key] = student
                done["students"] += 1
        else:
            if choice not in by_id:
                by_id[choice] = Student.objects.filter(pk=choice).first()
            student = by_id[choice]
            if not student:
                raise ValidationError(f"{r.line_label} : élève choisi introuvable.")
        invoice, was_created = add_prior_balance(student, school, r.start_year, r.balance)
        if was_created:
            done["invoices"] += 1
            done["total"] += r.balance
        else:
            done["already"] += 1
    return done
