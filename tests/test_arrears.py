"""V5-12 : reprise des soldes des années antérieures (Excel ou saisie), anciens élèves compris."""
import io
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook

from administration.models import Role, SchoolYear
from finance.models import FeeType, Invoice, Payment, PaymentBatch
from finance.services.arrears import add_prior_balance
from pedagogie.models import Classe, EducationLevel, Enrollment, Student


def classeur():
    """Même présentation que le tableau de l'établissement (Classeur1.xlsx)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Feuil1"
    ws.append(["Situation des paiements de frais de scolarité 2023 - 2024"])
    ws.append(["N°", "PRENOMS", "NOM", "Mt ANNUEL", "INSCRIPTION", "Oct", "Nov", "Déc", "Mt_perçu", "Reliquat"])
    ws.append([1, "Ali", "Diallo", 100000, 10000, 20000, 20000, None, 50000, 50000])
    ws.append([2, "Awa", "Keita", 100000, 10000, 30000, 30000, 30000, 100000, 0])          # soldée
    ws.append([3, "Moussa", "Traoré", 90000, 10000, 10000, None, None, 20000, 70000])     # parti
    ws.append([None, "TOTAL", None, 290000, None, None, None, None, 170000, 120000])
    ws.append([])
    ws.append(["Situation des paiements de frais de scolarité 2024 - 2025"])
    ws.append(["N°", "PRENOMS", "NOM", "Mt ANNUEL", "INSCRIPTION", "Oct", "Mt_perçu", "Reliquat"])
    ws.append([1, "Ali", "Diallo", 110000, 10000, 20000, 80000, 30000])
    ws.append([2, "Moussa", "TRAORE", 95000, 10000, 15000, 75000, 20000])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture(autouse=True)
def _tmp_data_dir(settings, tmp_path):
    settings.KIS_DATA_DIR = tmp_path


@pytest.fixture
def ctx(school):
    year = SchoolYear.objects.get(school=school)   # 2026-2027, année gérée dans KIS
    year.validate_catalog()
    classe = Classe(school=school, school_year=year, education_level=EducationLevel.objects.get(code="LYC-10"), group="A", name="")
    classe.save()
    scol = FeeType.objects.create(name="Scolarité", category="tuition", school_year=year, amount=Decimal("90000"), is_required=True)
    ali = Student.objects.create(first_name="Ali", last_name="Diallo", father_phone="76413630")
    current = Enrollment.objects.create(student=ali, classe=classe, school_year=year)
    return dict(school=school, year=year, ali=ali, current=current, scol=scol)


def _upload(client, data=None):
    return client.post("/finance/reprise-soldes/", {"file": SimpleUploadedFile("Classeur1.xlsx", data or classeur())})


def test_preview_then_import_is_idempotent(client, login, ctx):
    login(Role.ACCOUNTANT)
    r = _upload(client)
    report = r.context["report"]
    statuses = {(row.last_name, row.start_year): row.status for row in report.rows}
    assert statuses[("Diallo", 2023)] == "élève existant"
    assert statuses[("Keita", 2023)] == "soldé"
    assert statuses[("Traoré", 2023)] == "nouvel élève"
    assert ("TOTAL", 2023) not in statuses and all(row.first_name != "TOTAL" for row in report.rows)
    assert report.total == Decimal("170000")
    assert not Invoice.objects.filter(fee_type__category="arrears").exists()   # rien avant confirmation

    r = client.post("/finance/reprise-soldes/", {"token": r.context["token"], "confirm": "1"})
    assert r.status_code == 302
    arrears = Invoice.objects.filter(fee_type__category="arrears")
    assert arrears.count() == 4
    assert sum(i.total_amount for i in arrears) == Decimal("170000")
    archives = SchoolYear.objects.filter(is_archive=True)
    assert set(archives.values_list("name", flat=True)) == {"2023-2024", "2024-2025"}
    assert all(y.is_closed and not y.is_active for y in archives)
    moussa = Student.objects.get(last_name="Traoré")   # ancien élève créé une fois, retrouvé pour 2024
    assert moussa.enrollments.count() == 2
    # Réimport du même fichier : aucun doublon.
    r = _upload(client)
    assert r.context["report"].importable == []
    assert {row.status for row in r.context["report"].rows} <= {"déjà importé", "soldé"}


def test_arrears_paid_with_current_fees_on_one_receipt(client, login, ctx):
    login(Role.SECRETARY)
    old, _ = add_prior_balance(ctx["ali"], ctx["school"], 2024, Decimal("30000"))
    current = Invoice.objects.get(enrollment=ctx["current"], fee_type=ctx["scol"])
    page = client.get(f"/finance/encaissement/{ctx['current'].id}/")
    assert old in page.context["invoices"] and current in page.context["invoices"]
    r = client.post(f"/finance/encaissement/{ctx['current'].id}/", {
        "invoice": [old.id, current.id], f"amount_{old.id}": "30000", f"amount_{current.id}": "45000"})
    assert r.status_code == 302
    batch = PaymentBatch.objects.get()
    assert batch.total == Decimal("75000") and batch.payments.count() == 2
    old.refresh_from_db()
    assert old.reste == 0
    assert client.get(f"/finance/recus/{batch.id}/").content.startswith(b"%PDF")


def test_former_student_found_by_search_and_can_pay(client, login, ctx):
    login(Role.SECRETARY)
    former = Student.objects.create(first_name="Moussa", last_name="Traoré")
    inv, _ = add_prior_balance(former, ctx["school"], 2023, Decimal("70000"))
    r = client.get("/finance/payments/new/?q=Traor")
    assert [e.student for e in r.context["rows"]] == [former]
    enrollment = r.context["rows"][0]
    r = client.post(f"/finance/encaissement/{enrollment.id}/", {"invoice": [inv.id], f"amount_{inv.id}": "20000"})
    assert r.status_code == 302
    inv.refresh_from_db()
    assert inv.reste == Decimal("50000")
    page = client.get(f"/students/{former.id}/").content.decode()
    assert "Reprise des soldes" in page and "Encaisser" in page


def test_manual_prior_balance_and_rules(client, login, ctx):
    login(Role.ACCOUNTANT)
    url = f"/finance/eleves/{ctx['ali'].id}/reliquat/"
    r = client.post(url, {"start_year": "2025", "amount": "25 000"})
    assert r.status_code == 302
    assert Invoice.objects.get(fee_type__category="arrears").total_amount == Decimal("25000")
    r = client.post(url, {"start_year": "2025", "amount": "25000"})   # doublon refusé
    assert Invoice.objects.filter(fee_type__category="arrears").count() == 1
    r = client.post(url, {"start_year": "2026", "amount": "1000"})     # année gérée dans KIS
    assert r.status_code == 200
    with pytest.raises(Exception):
        add_prior_balance(ctx["ali"], ctx["school"], 2026, Decimal("1000"))


def test_archive_years_do_not_block_next_year_and_are_not_activable(ctx):
    add_prior_balance(ctx["ali"], ctx["school"], 2022, Decimal("1000"))
    archive = SchoolYear.objects.get(is_archive=True)
    assert archive.classes.get().display_name == "Reprise des soldes"
    # L'année gérée encore ouverte reste l'unique blocage pour créer la suivante.
    assert "2026-2027" in str(SchoolYear.creation_blocker(ctx["school"].id))


def test_unrecognised_file_gives_clear_message(client, login, ctx):
    login(Role.ACCOUNTANT)
    wb = Workbook(); wb.active.append(["Bonjour"]); buf = io.BytesIO(); wb.save(buf)
    r = _upload(client, buf.getvalue())
    msgs = [str(m) for m in r.context["messages"]]
    assert any("Aucun tableau reconnu" in m for m in msgs)
    r = client.post("/finance/reprise-soldes/", {"file": SimpleUploadedFile("x.xlsx", b"pas un excel")})
    assert any("lisible" in str(m) for m in r.context["messages"])


# --------------------------------------------------------------------------
# Retours V5-12 (suite) : présentation réelle de l'école, rattachement, modèle,
# saisie rapide, priorité de l'année en cours, reprise des versements.
# --------------------------------------------------------------------------
def classeur_reel():
    """Présentation du Classeur1 de l'école : en-tête d'école, blocs sans en-têtes,
    classe dans une colonne sans titre, ligne de totaux sans nom."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Pour_Kas"
    ws.append([None, None, "ECOLE FONDAMENTALE - JARDIN D'ENFANTS"])
    ws.append([None, None, None, "Tél.: 66 53 44 60"])
    ws.append([])
    ws.append(["Situation des paiements de frais de scolarité 2022 - 2023"])
    ws.append(["PRENOMS", "NOM", "Mt ANNUEL", "INSCRIPTION", "Octobre", "novembre", "Mt_perçu", "Reliquat"])
    ws.append(["Assitan", "MAGASSOUBA", 67500, 0, 10000, 30000, 40000, 27500])
    ws.append([None, None, None, None, None, None, 40000, 27500])            # totaux
    ws.append(["Situation des paiements de frais de scolarité 2023 - 2024"])  # pas d'en-têtes
    ws.append(["Seydou", "FANE", 90000, None, 0, 0, 0, 90000])
    ws.append(["Ba Aly", "FANE", 112500, None, 42500, 70000, 112500, 0])
    ws.append(["Situation des paiements de frais de scolarité 2025 - 2026"])
    ws.append(["Yacouba ", "DIABATÉ", 112500, None, 12500, 0, 50000, 62500])
    ws.append(["Adama F", "COULIBALY", 55000, None, 0, 50000, 50000, 5000])
    ws.append(["Ladji", "FANE", 90000, 0, 0, 0, 0, 90000, "GS"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_real_layout_blocks_without_headers_and_class_column(client, login, ctx):
    login(Role.ACCOUNTANT)
    Student.objects.create(first_name="Yacouba", last_name="DIABATE")
    adama = Student.objects.create(first_name="Adama Fousseyni", last_name="Coulibaly")
    report = _upload(client, classeur_reel()).context["report"]
    rows = {(r.last_name, r.first_name): r for r in report.rows}
    assert rows[("FANE", "Seydou")].start_year == 2023 and rows[("FANE", "Seydou")].balance == 90000
    assert rows[("FANE", "Ba Aly")].status == "soldé"
    assert rows[("FANE", "Ladji")].class_label == "GS"
    assert rows[("DIABATÉ", "Yacouba")].status == "élève existant"     # espaces et accents ignorés
    adama_row = rows[("COULIBALY", "Adama F")]
    assert adama_row.status == "à vérifier" and adama_row.default_choice == str(adama.pk)
    assert report.total == Decimal("27500") + 90000 + 62500 + 5000 + 90000


def test_manual_choices_are_applied(client, login, ctx):
    login(Role.ACCOUNTANT)
    adama = Student.objects.create(first_name="Adama Fousseyni", last_name="Coulibaly")
    ladji = Student.objects.create(first_name="Ladji Mamadou", last_name="Fane")
    r = _upload(client, classeur_reel())
    rows = {(row.last_name, row.first_name): row for row in r.context["report"].rows}
    post = {"token": r.context["token"], "confirm": "1",
            f"match_{rows[('COULIBALY', 'Adama F')].index}": "new",        # finalement un autre élève
            f"match_{rows[('FANE', 'Ladji')].index}": str(ladji.pk),
            f"match_{rows[('FANE', 'Seydou')].index}": "skip"}
    assert client.post("/finance/reprise-soldes/", post).status_code == 302
    assert not Invoice.objects.filter(enrollment__student=adama).exists()
    assert Invoice.objects.get(enrollment__student=ladji).total_amount == 90000
    assert Student.objects.filter(last_name="COULIBALY", first_name="Adama F").exists()
    assert not Student.objects.filter(first_name="Seydou").exists()


def test_homonyms_are_never_guessed(client, login, ctx):
    login(Role.ACCOUNTANT)
    Student.objects.create(first_name="Seydou", last_name="Fane")
    Student.objects.create(first_name="Seydou", last_name="FANE")
    r = _upload(client, classeur_reel())
    row = next(x for x in r.context["report"].rows if x.first_name == "Seydou")
    assert row.status == "à vérifier" and row.default_choice == "skip" and len(row.candidates) == 2


def test_kis_template_with_unspecified_year(client, login, ctx):
    login(Role.ACCOUNTANT)
    resp = client.get("/finance/reprise-soldes/modele.xlsx?eleves=1")
    assert resp.status_code == 200
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["Reliquats"]
    assert [c.value for c in ws[1]] == ["Matricule", "Prénom", "Nom", "Classe", "Année", "Reste dû"]
    assert ws.cell(2, 2).value == "Ali" and ws.cell(2, 4).value == "10ème année A"
    ws.cell(2, 5).value = "2024-2025"; ws.cell(2, 6).value = 30000
    ws.append(["", "Ali", "Diallo", "", "", 12000])                      # année inconnue
    buf = io.BytesIO(); wb.save(buf)
    r = _upload(client, buf.getvalue())
    report = r.context["report"]
    assert [(x.year_label, x.balance) for x in report.importable] == [("2024-2025", 30000), ("Non précisée", 12000)]
    client.post("/finance/reprise-soldes/", {"token": r.context["token"], "confirm": "1"})
    labels = set(Invoice.objects.filter(enrollment__student=ctx["ali"], fee_type__category="arrears").values_list("fee_type__name", flat=True))
    assert labels == {"Reliquat de scolarité 2024-2025", "Arriérés antérieurs (année non précisée)"}


def test_quick_entry_by_class(client, login, ctx):
    login(Role.ACCOUNTANT)
    classe = ctx["current"].classe
    binta = Student.objects.create(first_name="Binta", last_name="Keita")
    Enrollment.objects.create(student=binta, classe=classe, school_year=ctx["year"])
    url = f"/finance/reprise-soldes/saisie-rapide/?classe={classe.id}&annee=2025"
    page = client.get(url)
    assert len(page.context["rows"]) == 2
    r = client.post("/finance/reprise-soldes/saisie-rapide/", {"classe": classe.id, "annee": "2025",
                                                              f"amount_{ctx['ali'].id}": "15 000", f"amount_{binta.id}": ""})
    assert r.status_code == 302
    assert Invoice.objects.get(fee_type__category="arrears").total_amount == 15000
    # Déjà saisi : affiché, pas de seconde saisie possible.
    page = client.get(url).content.decode()
    assert "déjà enregistré" in page and f"amount_{ctx['ali'].id}" not in page
    r = client.post("/finance/reprise-soldes/saisie-rapide/", {"classe": classe.id, "annee": "nc", f"amount_{binta.id}": "8000"})
    assert Invoice.objects.get(enrollment__student=binta, fee_type__category="arrears").fee_type.name == "Arriérés antérieurs (année non précisée)"


def test_current_year_first_then_debts_oldest_first(client, login, ctx):
    login(Role.SECRETARY)
    add_prior_balance(ctx["ali"], ctx["school"], 2024, Decimal("30000"))
    add_prior_balance(ctx["ali"], ctx["school"], 2022, Decimal("20000"))
    add_prior_balance(ctx["ali"], ctx["school"], None, Decimal("5000"))
    page = client.get(f"/finance/encaissement/{ctx['current'].id}/")
    names = [i.fee_type.name for i in page.context["invoices"]]
    assert names == ["Scolarité", "Arriérés antérieurs (année non précisée)",
                     "Reliquat de scolarité 2022-2023", "Reliquat de scolarité 2024-2025"]
    assert "Répartir automatiquement" in page.content.decode()


def test_carryover_payment_without_receipt(client, login, ctx):
    current = Invoice.objects.get(enrollment=ctx["current"], fee_type=ctx["scol"])
    login(Role.SECRETARY)                     # la caisse ne peut pas faire de reprise
    client.post(f"/finance/encaissement/{ctx['current'].id}/", {
        "invoice": [current.id], f"amount_{current.id}": "40000", "carryover": "1", "carryover_date": "2026-10-02"})
    assert PaymentBatch.objects.count() == 1 and not Payment.objects.filter(is_carryover=True).exists()
    Payment.objects.all().delete(); PaymentBatch.objects.all().delete(); current.refresh_from_db()
    current.paid_amount = 0; current.save()
    client.logout()
    login(Role.ACCOUNTANT, username="compta")
    r = client.post(f"/finance/encaissement/{ctx['current'].id}/", {
        "invoice": [current.id], f"amount_{current.id}": "40000", "carryover": "1", "carryover_date": "2026-10-02"})
    assert r.status_code == 302
    p = Payment.objects.get()
    assert p.is_carryover and p.batch is None and p.receipt_number.startswith("REPRISE-")
    assert str(p.date) == "2026-10-02"
    current.refresh_from_db()
    assert current.reste == Decimal("50000")
    assert client.get(f"/finance/payments/{p.id}/receipt/").status_code == 404
    # La numérotation des reçus n'est pas touchée.
    from common.services.numbering_service import NumberingService
    assert NumberingService.next_receipt_number().endswith("00001")
    r = client.post(f"/finance/encaissement/{ctx['current'].id}/", {
        "invoice": [current.id], f"amount_{current.id}": "1000", "carryover": "1", "carryover_date": "2999-01-01"})
    assert r.status_code == 200 and Payment.objects.count() == 1
